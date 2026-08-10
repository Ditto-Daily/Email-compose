"""Manage the editable response-template library and Excel interchange."""

from __future__ import annotations

import json
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from config import STANDARD_TEMPLATE_PATH, TEMPLATES_JSON_PATH


def _write_text_atomic(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(f"{path.suffix}.tmp")
    temporary.write_text(content, encoding="utf-8")
    temporary.replace(path)


def _load_payload() -> dict[str, Any]:
    try:
        data = json.loads(TEMPLATES_JSON_PATH.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except (FileNotFoundError, json.JSONDecodeError):
        return {}


def load_templates() -> list[dict[str, Any]]:
    """Load templates, falling back to the legacy standard.txt content."""
    data = _load_payload()
    if data:
        templates = data.get("templates", [])
        if isinstance(templates, list):
            active_categories = data.get("active_categories")
            if isinstance(active_categories, list):
                active_set = set(active_categories)
                for template in templates:
                    template["active"] = template.get("category", "General") in active_set
            return templates

    legacy = ""
    try:
        legacy = STANDARD_TEMPLATE_PATH.read_text(encoding="utf-8").strip()
    except FileNotFoundError:
        pass
    if not legacy:
        return []
    return [
        {
            "id": str(uuid.uuid4()),
            "name": "Standard response",
            "category": "General",
            "content": legacy,
            "active": True,
        }
    ]


def load_category_settings(
    templates: list[dict[str, Any]] | None = None,
) -> tuple[list[str], list[str]]:
    """Return all categories and the subset currently available to Gemini."""
    data = _load_payload()
    stored_categories = data.get("categories")
    stored_active = data.get("active_categories")
    if isinstance(stored_categories, list) and isinstance(stored_active, list):
        categories = [str(category) for category in stored_categories]
        active = [str(category) for category in stored_active if category in categories]
        return categories, active

    templates = templates if templates is not None else load_templates()
    categories: list[str] = []
    active: list[str] = []
    for template in templates:
        category = str(template.get("category") or "General").strip()
        if category not in categories:
            categories.append(category)
        if template.get("active", True) and category not in active:
            active.append(category)
    return categories, active


def active_template_context(templates: list[dict[str, Any]] | None = None) -> str:
    """Render all active templates as structured Gemini context."""
    templates = templates if templates is not None else load_templates()
    active = [template for template in templates if template.get("active", True)]
    if not active:
        return "No active response templates."

    sections: list[str] = ["# Active Customer Success Response Library"]
    current_category = None
    for template in sorted(
        active,
        key=lambda item: (
            str(item.get("category", "General")).lower(),
            str(item.get("name", "")).lower(),
        ),
    ):
        category = str(template.get("category") or "General").strip()
        if category != current_category:
            sections.extend(["", f"## {category}"])
            current_category = category
        sections.extend(
            [
                "",
                f"### {str(template.get('name') or 'Untitled').strip()}",
                str(template.get("content") or "").strip(),
            ]
        )
    return "\n".join(sections).rstrip() + "\n"


def save_templates(
    templates: list[dict[str, Any]],
    categories: list[str] | None = None,
    active_categories: list[str] | None = None,
) -> None:
    """Persist category settings, templates, and Gemini's text context."""
    saved_categories, saved_active = load_category_settings(templates)
    categories = list(categories) if categories is not None else saved_categories
    active_categories = (
        list(active_categories) if active_categories is not None else saved_active
    )

    for template in templates:
        category = str(template.get("category") or "General").strip()
        if category not in categories:
            categories.append(category)
    categories = list(dict.fromkeys(category.strip() for category in categories if category.strip()))
    active_categories = [
        category
        for category in dict.fromkeys(active_categories)
        if category in categories
    ]
    active_set = set(active_categories)

    normalized = []
    for template in templates:
        name = str(template.get("name", "")).strip()
        content = str(template.get("content", "")).strip()
        category = str(template.get("category") or "General").strip()
        if not name or not content:
            continue
        normalized.append(
            {
                "id": str(template.get("id") or uuid.uuid4()),
                "name": name,
                "category": category,
                "content": content,
                "active": category in active_set,
            }
        )

    payload = {
        "version": 2,
        "categories": categories,
        "active_categories": active_categories,
        "templates": normalized,
    }
    _write_text_atomic(
        TEMPLATES_JSON_PATH,
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
    )
    _write_text_atomic(STANDARD_TEMPLATE_PATH, active_template_context(normalized))


def import_excel(path: str | Path) -> list[dict[str, Any]]:
    """Convert the first worksheet's Q&A rows into grouped templates."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    templates: list[dict[str, Any]] = []
    category = "General"

    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), 1):
        values = [str(value).strip() for value in row if value is not None and str(value).strip()]
        if not values:
            continue
        if row_number == 1 and values[0].lower() in {"question", "name", "template"}:
            continue

        name = values[0]
        answer_parts = values[1:]
        if not answer_parts:
            category = name
            continue

        templates.append(
            {
                "id": str(uuid.uuid4()),
                "name": name,
                "category": category,
                "content": "\n\n".join(answer_parts),
                "active": True,
            }
        )
    workbook.close()
    return templates


def export_excel(templates: list[dict[str, Any]]) -> bytes:
    """Return an Excel workbook containing the complete template library."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Templates"
    worksheet.append(
        [
            "Category",
            "Template / Question",
            "Response",
            "Category available to Gemini",
        ]
    )
    for template in templates:
        worksheet.append(
            [
                template.get("category", "General"),
                template.get("name", ""),
                template.get("content", ""),
                "Yes" if template.get("active", True) else "No",
            ]
        )
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 48
    worksheet.column_dimensions["C"].width = 100
    worksheet.column_dimensions["D"].width = 12

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
